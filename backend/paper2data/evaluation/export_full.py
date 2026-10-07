"""Полный экспорт корпуса: CSV/XLSX, JSON, GraphML, SDF (опционально RDKit)."""

from __future__ import annotations

import csv
import json
import xml.etree.ElementTree as ET
from collections import defaultdict
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from sqlalchemy import select
from sqlalchemy.orm import Session

from paper2data.core.project import ProjectState
from paper2data.db.models import Document
from paper2data.graph import service as graph_service


def _stamp() -> str:
    return datetime.now(timezone.utc).strftime("%Y%m%d_%H%M%S")


def _load_project_graph(
    session: Session,
    document_ids: list[int] | None = None,
) -> dict[str, Any]:
    stmt = select(Document).order_by(Document.id)
    if document_ids:
        stmt = stmt.where(Document.id.in_(document_ids))
    docs = session.scalars(stmt).all()
    all_nodes: list[dict[str, Any]] = []
    all_edges: list[dict[str, Any]] = []
    documents: list[dict[str, Any]] = []
    for doc in docs:
        g = graph_service.get_graph(session, doc.id)
        documents.append(
            {
                "id": doc.id,
                "title": doc.title,
                "doi": doc.doi,
                "status": doc.status,
            }
        )
        all_nodes.extend(g["nodes"])
        all_edges.extend(g["edges"])
    return {"documents": documents, "nodes": all_nodes, "edges": all_edges}


def _write_csv_by_type(exports: Path, base: str, nodes: list[dict[str, Any]]) -> list[str]:
    by_type: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for node in nodes:
        by_type[str(node.get("entity_type") or "Unknown")].append(node)
    paths: list[str] = []
    for type_id, group in by_type.items():
        attr_keys: list[str] = []
        seen: set[str] = set()
        for n in group:
            for k in (n.get("attributes") or {}).keys():
                if k not in seen:
                    seen.add(k)
                    attr_keys.append(k)
        path = exports / f"{base}_{type_id}.csv"
        with path.open("w", encoding="utf-8-sig", newline="") as fh:
            writer = csv.writer(fh)
            writer.writerow(
                [
                    "node_id",
                    "document_id",
                    "label",
                    "status",
                    "entity_type",
                    "canonical_id",
                    *attr_keys,
                ]
            )
            for n in group:
                attrs = n.get("attributes") or {}
                writer.writerow(
                    [
                        n.get("id"),
                        n.get("document_id"),
                        n.get("label") or "",
                        n.get("status") or "",
                        type_id,
                        n.get("canonical_id") or "",
                        *[attrs.get(k, "") for k in attr_keys],
                    ]
                )
        paths.append(str(path))
    return paths


def _write_xlsx(exports: Path, base: str, nodes: list[dict[str, Any]], edges: list[dict[str, Any]]) -> str | None:
    try:
        from openpyxl import Workbook
    except ImportError:
        return None

    wb = Workbook()
    # nodes by type
    by_type: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for node in nodes:
        by_type[str(node.get("entity_type") or "Unknown")].append(node)

    first = True
    for type_id, group in by_type.items():
        title = type_id[:31] or "Unknown"
        ws = wb.active if first else wb.create_sheet(title)
        if first:
            ws.title = title
            first = False
        attr_keys: list[str] = []
        seen: set[str] = set()
        for n in group:
            for k in (n.get("attributes") or {}).keys():
                if k not in seen:
                    seen.add(k)
                    attr_keys.append(k)
        headers = [
            "node_id",
            "document_id",
            "label",
            "status",
            "canonical_id",
            *attr_keys,
        ]
        ws.append(headers)
        for n in group:
            attrs = n.get("attributes") or {}
            ws.append(
                [
                    n.get("id"),
                    n.get("document_id"),
                    n.get("label") or "",
                    n.get("status") or "",
                    n.get("canonical_id") or "",
                    *[attrs.get(k, "") for k in attr_keys],
                ]
            )

    ws_e = wb.create_sheet("edges")
    ws_e.append(
        [
            "edge_id",
            "document_id",
            "relation_type",
            "source_node_id",
            "target_node_id",
            "status",
            "attributes_json",
        ]
    )
    for e in edges:
        ws_e.append(
            [
                e.get("id"),
                e.get("document_id"),
                e.get("relation_type"),
                e.get("source_node_id"),
                e.get("target_node_id"),
                e.get("status"),
                json.dumps(e.get("attributes") or {}, ensure_ascii=False),
            ]
        )

    path = exports / f"{base}.xlsx"
    wb.save(path)
    return str(path)


def _write_graphml(exports: Path, base: str, nodes: list[dict[str, Any]], edges: list[dict[str, Any]]) -> str:
    """GraphML без обязательного NetworkX."""
    ns = "http://graphml.graphdrawing.org/xmlns"
    ET.register_namespace("", ns)
    root = ET.Element(f"{{{ns}}}graphml")
    for key_id, attr_name, for_what, attr_type in [
        ("d_label", "label", "node", "string"),
        ("d_type", "entity_type", "node", "string"),
        ("d_doc", "document_id", "node", "string"),
        ("d_canon", "canonical_id", "node", "string"),
        ("d_rel", "relation_type", "edge", "string"),
    ]:
        k = ET.SubElement(root, f"{{{ns}}}key")
        k.set("id", key_id)
        k.set("for", for_what)
        k.set("attr.name", attr_name)
        k.set("attr.type", attr_type)

    graph = ET.SubElement(root, f"{{{ns}}}graph")
    graph.set("id", "G")
    graph.set("edgedefault", "directed")

    for n in nodes:
        node_el = ET.SubElement(graph, f"{{{ns}}}node")
        node_el.set("id", f"n{n['id']}")
        for key_id, val in [
            ("d_label", n.get("label") or ""),
            ("d_type", n.get("entity_type") or ""),
            ("d_doc", str(n.get("document_id") or "")),
            ("d_canon", n.get("canonical_id") or ""),
        ]:
            data = ET.SubElement(node_el, f"{{{ns}}}data")
            data.set("key", key_id)
            data.text = str(val)

    for e in edges:
        edge_el = ET.SubElement(graph, f"{{{ns}}}edge")
        edge_el.set("id", f"e{e['id']}")
        edge_el.set("source", f"n{e['source_node_id']}")
        edge_el.set("target", f"n{e['target_node_id']}")
        data = ET.SubElement(edge_el, f"{{{ns}}}data")
        data.set("key", "d_rel")
        data.text = str(e.get("relation_type") or "")

    path = exports / f"{base}.graphml"
    tree = ET.ElementTree(root)
    tree.write(path, encoding="utf-8", xml_declaration=True)
    return str(path)


def _try_write_sdf(exports: Path, base: str, nodes: list[dict[str, Any]]) -> dict[str, Any]:
    """SDF через RDKit, если доступен; иначе skip + note."""
    note = None
    try:
        from rdkit import Chem
        from rdkit.Chem import SDWriter
    except ImportError:
        return {
            "sdf_path": None,
            "skipped": True,
            "note": "RDKit не установлен — SDF пропущен. Установите rdkit для экспорта структур.",
            "molecules": 0,
        }

    compounds = [
        n
        for n in nodes
        if (n.get("entity_type") or "").lower() in {"compound", "component"}
        and (n.get("attributes") or {}).get("smiles")
    ]
    if not compounds:
        return {
            "sdf_path": None,
            "skipped": True,
            "note": "Нет узлов Compound/Component с SMILES.",
            "molecules": 0,
        }

    path = exports / f"{base}.sdf"
    writer = SDWriter(str(path))
    written = 0
    errors = 0
    for n in compounds:
        smiles = str((n.get("attributes") or {}).get("smiles") or "").strip()
        mol = Chem.MolFromSmiles(smiles)
        if mol is None:
            errors += 1
            continue
        mol.SetProp("_Name", str(n.get("label") or n.get("id")))
        mol.SetProp("node_id", str(n.get("id")))
        mol.SetProp("document_id", str(n.get("document_id") or ""))
        if n.get("canonical_id"):
            mol.SetProp("canonical_id", str(n["canonical_id"]))
        writer.write(mol)
        written += 1
    writer.close()
    if written == 0:
        path.unlink(missing_ok=True)
        return {
            "sdf_path": None,
            "skipped": True,
            "note": f"RDKit не смог разобрать SMILES ({errors} ошибок).",
            "molecules": 0,
        }
    return {
        "sdf_path": str(path),
        "skipped": False,
        "note": note,
        "molecules": written,
        "errors": errors,
    }


def export_full(
    session: Session,
    project: ProjectState,
    *,
    document_ids: list[int] | None = None,
    include_sdf: bool = True,
) -> dict[str, Any]:
    exports = project.exports_dir
    exports.mkdir(parents=True, exist_ok=True)
    base = f"corpus_{_stamp()}"
    graph = _load_project_graph(session, document_ids)
    nodes = graph["nodes"]
    edges = graph["edges"]

    json_path = exports / f"{base}_graph.json"
    with json_path.open("w", encoding="utf-8") as fh:
        json.dump(graph, fh, ensure_ascii=False, indent=2)

    csv_paths = _write_csv_by_type(exports, base, nodes)
    xlsx_path = _write_xlsx(exports, base, nodes, edges)
    graphml_path = _write_graphml(exports, base, nodes, edges)
    sdf_info: dict[str, Any] = {"skipped": True, "note": "SDF отключён", "molecules": 0}
    if include_sdf:
        sdf_info = _try_write_sdf(exports, base, nodes)

    return {
        "exports_dir": str(exports),
        "base": base,
        "json_path": str(json_path),
        "csv_paths": csv_paths,
        "xlsx_path": xlsx_path,
        "xlsx_note": None
        if xlsx_path
        else "openpyxl не установлен — XLSX пропущен. pip install openpyxl",
        "graphml_path": graphml_path,
        "sdf": sdf_info,
        "document_count": len(graph["documents"]),
        "node_count": len(nodes),
        "edge_count": len(edges),
    }
